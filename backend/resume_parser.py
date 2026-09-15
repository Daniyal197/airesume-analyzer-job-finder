import io
import re
from pathlib import Path

import pdfplumber
from docx import Document


class UnsupportedFileTypeError(Exception):
    pass 
class EmptyResumeError(Exception):
    pass

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_text_from_pdf(file_bytes: bytes) -> str:
    text_chunks = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_chunks.append(page_text)
    return "\n".join(text_chunks)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))
    text_chunks = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text_chunks.append(cell.text)

    return "\n".join(text_chunks)


def normalize_text(raw_text: str) -> str:
    text = re.sub(r"[•●▪◦‣]", "-", raw_text)

    text = re.sub(r"\n{3,}", "\n\n", text)

    lines = [line.rstrip() for line in text.split("\n")]
    while lines and not lines[0].strip():
        lines.pop(0)
    while lines and not lines[-1].strip():
        lines.pop()

    return "\n".join(lines).strip()


def parse_resume(filename: str, file_bytes: bytes) -> dict:
    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            f"'{extension}' is not supported. Upload a PDF or DOCX resume."
        )

    if extension == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    else:
        raw_text = extract_text_from_docx(file_bytes)

    normalized = normalize_text(raw_text)

    if len(normalized) < 20:
        raise EmptyResumeError(
            "Couldn't find readable text in this file. "
            "If it's a scanned/image-based PDF, OCR isn't supported yet."
        )

    return {
        "filename": filename,
        "file_type": extension,
        "char_count": len(normalized),
        "word_count": len(normalized.split()),
        "text": normalized,
    }